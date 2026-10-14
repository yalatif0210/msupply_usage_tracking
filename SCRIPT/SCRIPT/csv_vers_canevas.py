"""Assemble les CSV produits par extraction_msupply.sql dans une copie du
gabarit Canevas_Extraction_mSupply_DQA_v2.xlsx, puis vérifie le classeur obtenu.

Usage :
    python csv_vers_canevas.py <dossier_csv> [gabarit.xlsx] [sortie.xlsx]
    python csv_vers_canevas.py --verifier <classeur.xlsx>

Par défaut : gabarit = Canevas_Extraction_mSupply_DQA_v2.xlsx à côté de ce
script, sortie = <dossier_csv>/Canevas_Extraction_mSupply_DQA_v2_<date>.xlsx
(date = date_extraction lue dans metadata.csv).

Règles appliquées :
  - ligne 5 (en-têtes) jamais modifiée ; les données s'écrivent à partir de
    la ligne 6. La ligne d'exemple du gabarit est REMPLACÉE par la première
    donnée réelle : il n'y a rien à supprimer ensuite ;
  - chaque cellule écrite reçoit une police normale. Sans cela, la première
    ligne de données garde le gris italique de l'exemple, passe pour
    l'exemple et se fait supprimer à la main (lot du 2026-10-12 : 8 onglets
    rejetés pour une ligne de moins que le compteur METADATA) ;
  - colonnes repérées par leur NOM en ligne 5 (l'ordre du CSV n'importe pas) ;
  - toutes les valeurs sont écrites en TEXTE (format « @ ») : les dates
    restent AAAA-MM-JJ, jamais un numéro de série Excel. Un texte qui commence
    par « = » reste du texte (openpyxl en ferait une formule, vide à la
    lecture). Les caractères de contrôle interdits en XML, que l'on trouve
    parfois dans les commentaires mSupply, sont remplacés par une espace
    (openpyxl refuserait sinon d'écrire la cellule) ;
  - METADATA : seule la colonne B « Valeur » est remplie, en face du champ ;
  - MODE_EMPLOI : dans le classeur produit, les consignes « supprimer la ligne
    d'exemple » (valables pour une saisie à la main) sont remplacées par
    « ne rien supprimer » ;
  - vérification finale : le classeur enregistré est relu ; chaque compteur
    nb_lignes_* de METADATA doit égaler le nombre de lignes de données de son
    onglet, et les champs METADATA obligatoires doivent être renseignés. Tout
    écart fait échouer le script (code de sortie 1).

--verifier rejoue cette vérification sur un classeur existant : à lancer avant
l'envoi si le fichier a été ouvert ou réenregistré dans Excel entre-temps.
"""
import csv
import re
import shutil
import sys
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font

# Sous Windows, quand la sortie passe par Git Bash, Python écrit en cp1252
# alors que le terminal attend de l'UTF-8 : accents illisibles, et plantage
# sur tout caractère absent de cp1252. Sortie forcée en UTF-8, sans plantage.
for _flux in (sys.stdout, sys.stderr):
    try:
        _flux.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

ONGLETS = {
    "SITES": "sites.csv",
    "SITE_PROGRAMME": "site_programme.csv",
    "PRODUITS": "produits.csv",
    "TRANSACTIONS": "transactions.csv",
    "LIGNES_TRANSACTIONS": "lignes_transactions.csv",
    "STOCK_ACTUEL": "stock_actuel.csv",
    "INVENTAIRES": "inventaires.csv",
    "INVENTAIRE_LIGNES": "inventaire_lignes.csv",
    "REQUISITIONS": "requisitions.csv",
}
LIGNE_ENTETE = 5
PREMIERE_LIGNE = 6
METADATA_OBLIGATOIRES = ("date_extraction", "cycle_paqd", "mois_paqd", "mois_debut", "mois_fin",
                         "fenetre_debut", "fenetre_fin")

CONSIGNE_SANS_SUPPRESSION = (
    "Classeur produit par lancer_extraction.sh : la ligne 6 de chaque onglet contient "
    "déjà la première donnée réelle (la ligne d'exemple a été remplacée). Ne rien "
    "supprimer. Envoyer le fichier tel quel, sans le réenregistrer dans Excel ; s'il a "
    "été ouvert, le contrôler avec : python csv_vers_canevas.py --verifier <fichier>."
)
# Une consigne de suppression de l'exemple peut tenir en une phrase
# (« Supprimer les 10 lignes d'exemple ») ou en deux (« ... est un exemple.
# La supprimer avant l'envoi. ») : on repère les deux mots dans la cellule.
RE_EXEMPLE = re.compile(r"\bexemple\b", re.I)
RE_SUPPRIMER = re.compile(r"\bsupprim", re.I)
# Mention restée exacte mais trompeuse une fois l'exemple remplacé.
HORS_EXEMPLE = ("(hors en-tête, hors exemple)", "(lignes de données, à partir de la ligne 6)")


def lire_csv(chemin):
    with open(chemin, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def nettoyer(valeur):
    """Valeur à écrire : None si vide ; caractères interdits en XML remplacés."""
    if valeur in (None, ""):
        return None, False
    propre = ILLEGAL_CHARACTERS_RE.sub(" ", valeur)
    return propre, propre != valeur


def police_donnee(ws, colonne):
    """Police de l'en-tête (même famille, même taille), sans gras, italique ni couleur."""
    entete = ws.cell(LIGNE_ENTETE, colonne).font
    return Font(name=entete.name, size=entete.size)


def remplir_onglet(ws, lignes):
    entetes = {}
    for c in range(1, ws.max_column + 1):
        nom = ws.cell(LIGNE_ENTETE, c).value
        if nom:
            entetes[nom] = c
    if lignes:
        inconnues = set(lignes[0]) - set(entetes)
        if inconnues:
            print(f"  ! {ws.title} : colonnes CSV absentes du gabarit, ignorées : {sorted(inconnues)}")
    manquantes = [n for n in entetes if not lignes or n not in lignes[0]]
    polices = {c: police_donnee(ws, c) for c in entetes.values()}
    # Effacer la ligne d'exemple (et tout résidu) avant d'écrire, valeur ET
    # mise en forme : une ligne effacée ne doit plus ressembler à l'exemple.
    for r in range(PREMIERE_LIGNE, ws.max_row + 1):
        for c in entetes.values():
            ws.cell(r, c).value = None
            ws.cell(r, c).font = polices[c]
    nettoyees = 0
    for i, ligne in enumerate(lignes):
        r = PREMIERE_LIGNE + i
        for nom, c in entetes.items():
            valeur, modifiee = nettoyer(ligne.get(nom))
            nettoyees += modifiee
            cellule = ws.cell(r, c)
            cellule.number_format = "@"
            cellule.font = polices[c]
            cellule.value = valeur
            if valeur is not None and valeur.startswith("="):
                cellule.data_type = "s"  # texte, pas une formule
    return manquantes, nettoyees


def adapter_mode_emploi(wb):
    if "MODE_EMPLOI" not in wb.sheetnames:
        return 0
    n = 0
    for ligne in wb["MODE_EMPLOI"].iter_rows():
        for cellule in ligne:
            texte = cellule.value
            if not isinstance(texte, str):
                continue
            if RE_EXEMPLE.search(texte) and RE_SUPPRIMER.search(texte):
                numero = re.match(r"^\s*\d+\.\s*", texte)
                cellule.value = (numero.group(0) if numero else "") + CONSIGNE_SANS_SUPPRESSION
                n += 1
            elif HORS_EXEMPLE[0] in texte:
                cellule.value = texte.replace(*HORS_EXEMPLE)
    return n


def compter_lignes(ws):
    """Lignes de données (à partir de la ligne 6) ayant au moins une valeur dans
    une colonne nommée en ligne 5 : la règle de comptage du tableau de bord."""
    colonnes = [c for c in range(1, ws.max_column + 1) if ws.cell(LIGNE_ENTETE, c).value]
    n = 0
    for r in range(PREMIERE_LIGNE, ws.max_row + 1):
        if any(str(ws.cell(r, c).value).strip() not in ("", "None") for c in colonnes):
            n += 1
    return n


def verifier(chemin):
    """Erreurs bloquantes du classeur (liste vide si conforme)."""
    wb = openpyxl.load_workbook(chemin)
    erreurs = []
    if "METADATA" not in wb.sheetnames:
        return ["onglet METADATA absent"]
    ws = wb["METADATA"]
    metadata = {}
    for r in range(1, ws.max_row + 1):
        champ = ws.cell(r, 1).value
        if isinstance(champ, str) and champ.strip():
            metadata[champ.strip()] = ws.cell(r, 2).value
    for champ in METADATA_OBLIGATOIRES:
        if str(metadata.get(champ) or "").strip() == "":
            erreurs.append(f"METADATA : champ « {champ} » vide")
    for onglet in ONGLETS:
        if onglet not in wb.sheetnames:
            erreurs.append(f"onglet {onglet} absent")
            continue
        reel = compter_lignes(wb[onglet])
        declare = str(metadata.get(f"nb_lignes_{onglet}") or "").strip()
        if not declare.isdigit():
            erreurs.append(f"{onglet} : compteur nb_lignes_{onglet} absent ou non numérique ({declare!r})")
        elif int(declare) != reel:
            erreurs.append(f"{onglet} : METADATA annonce {declare} lignes, le classeur en contient {reel}")
    return erreurs


def rapporter(chemin, erreurs):
    if erreurs:
        print(f"\nVÉRIFICATION ÉCHOUÉE : {chemin}")
        for e in erreurs:
            print(f"  ! {e}")
        print("  -> ce classeur serait rejeté au dépôt. Relancer lancer_extraction.sh "
              "et envoyer le fichier produit sans le retoucher.")
        return 1
    print(f"\nVérification OK : {chemin}")
    print("  compteurs METADATA = lignes réelles sur les 9 onglets ; METADATA complète.")
    return 0


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    if sys.argv[1] == "--verifier":
        if len(sys.argv) < 3:
            sys.exit(__doc__)
        sys.exit(rapporter(sys.argv[2], verifier(sys.argv[2])))

    dossier = Path(sys.argv[1])
    gabarit = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(__file__).with_name("Canevas_Extraction_mSupply_DQA_v2.xlsx")

    metadata = {l["champ"]: l["valeur"] for l in lire_csv(dossier / "metadata.csv")}
    date = metadata.get("date_extraction", "sans_date")
    sortie = Path(sys.argv[3]) if len(sys.argv) > 3 else dossier / f"Canevas_Extraction_mSupply_DQA_v2_{date}.xlsx"

    shutil.copyfile(gabarit, sortie)
    wb = openpyxl.load_workbook(sortie)

    for onglet, fichier in ONGLETS.items():
        lignes = lire_csv(dossier / fichier)
        manquantes, nettoyees = remplir_onglet(wb[onglet], lignes)
        print(f"{onglet:<20} {len(lignes):>6} lignes")
        if lignes and manquantes:
            print(f"  ! {onglet} : colonnes du gabarit sans donnée dans le CSV : {manquantes}")
        if nettoyees:
            print(f"  info : {onglet} : {nettoyees} valeur(s) contenaient des caractères de contrôle, remplacés par une espace")
        attendu = metadata.get(f"nb_lignes_{onglet}")
        if attendu is not None and int(attendu) != len(lignes):
            print(f"  ! {onglet} : METADATA annonce {attendu} lignes, CSV en contient {len(lignes)}")

    ws = wb["METADATA"]
    for r in range(LIGNE_ENTETE, ws.max_row + 1):
        champ = ws.cell(r, 1).value
        if champ in metadata:
            ws.cell(r, 2).number_format = "@"
            ws.cell(r, 2).value = metadata[champ] or None
        elif champ:
            print(f"  ! METADATA : champ « {champ} » absent de metadata.csv")

    adapter_mode_emploi(wb)
    wb.save(sortie)
    print(f"\nClasseur écrit : {sortie}")
    sys.exit(rapporter(sortie, verifier(sortie)))


if __name__ == "__main__":
    main()
